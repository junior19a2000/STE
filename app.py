import marimo

__generated_with = "0.13.8"
app = marimo.App(width="full")


@app.cell(hide_code=True)
def _():
    import marimo as mo
    import zipfile
    import requests
    import numpy as np
    import pandas as pd
    import altair as alt
    import re
    import io
    from io import BytesIO
    from datetime import datetime
    from datetime import date
    from docxtpl import DocxTemplate, InlineImage
    from docx.shared import Cm 
    return (
        BytesIO,
        Cm,
        DocxTemplate,
        InlineImage,
        alt,
        datetime,
        io,
        mo,
        np,
        pd,
        re,
        requests,
        zipfile,
    )


@app.cell(hide_code=True)
def _():
    departamentos_region = {
        "AMAZONAS": "AMAZONAS",
        "ANCASH": "ANCASH",
        "APURIMAC": "APURIMAC",
        "AREQUIPA": "AREQUIPA",
        "AYACUCHO": "AYACUCHO",
        "CAJAMARCA": "CAJAMARCA",
        "CUSCO": "CUSCO",
        "HUANCAVELICA": "HUANCAVELICA",
        "HUANUCO": "HUANUCO",
        "ICA": "ICA",
        "JUNIN": "JUNIN",
        "LA LIBERTAD": "LA LIBERTAD",
        "LAMBAYEQUE": "LAMBAYEQUE",
        "LORETO": "LORETO",
        "MADRE DE DIOS": "MADRE DE DIOS",
        "MOQUEGUA": "MOQUEGUA",
        "PASCO": "PASCO",
        "PIURA": "PIURA",
        "PROV. CONST. DEL CALLAO": "LIMA NORTE",
        "PUNO": "PUNO",
        "SAN MARTIN": "SAN MARTIN",
        "TACNA": "TACNA",
        "TUMBES": "TUMBES",
        "UCAYALI": "UCAYALI"
    }
    provincias_region = {
        "BARRANCA": "LIMA NORTE",
        "HUAURA": "LIMA NORTE",
        "CAJATAMBO": "LIMA NORTE",
        "OYON": "LIMA NORTE",
        "HUARAL": "LIMA NORTE",
        "CANTA": "LIMA NORTE",
        "CAÑETE": "LIMA SUR",
        "HUAROCHIRI": "LIMA SUR",
        "YAUYOS": "LIMA SUR"
    }
    distritos_region = {
        "BELLAVISTA": "LIMA NORTE",
        "CALLAO": "LIMA NORTE",
        "PROV. CONST. DEL CALLAO": "LIMA NORTE",
        "CARMEN DE LA LEGUA REYNOSO": "LIMA NORTE",
        "LA PUNTA": "LIMA NORTE",
        "MI PERU": "LIMA NORTE",
        "LA PERLA": "LIMA NORTE",
        "VENTANILLA": "LIMA NORTE",
        "BREÑA": "LIMA NORTE",
        "JESUS MARIA": "LIMA NORTE",
        "LIMA": "LIMA NORTE",
        "LINCE": "LIMA NORTE",
        "MAGDALENA DEL MAR": "LIMA NORTE",
        "PUEBLO LIBRE": "LIMA NORTE",
        "RIMAC": "LIMA NORTE",
        "SAN ISIDRO": "LIMA NORTE",
        "SAN MIGUEL": "LIMA NORTE",
        "EL AGUSTINO": "LIMA NORTE",
        "SAN JUAN DE LURIGANCHO": "LIMA NORTE",
        "SANTA ANITA": "LIMA SUR",
        "ANCON": "LIMA NORTE",
        "CARABAYLLO": "LIMA NORTE",
        "COMAS": "LIMA NORTE",
        "INDEPENDENCIA": "LIMA NORTE",
        "LOS OLIVOS": "LIMA NORTE",
        "PUENTE PIEDRA": "LIMA NORTE",
        "SAN MARTIN DE PORRES": "LIMA NORTE",
        "SANTA ROSA": "LIMA NORTE",
        "BARRANCO": "LIMA SUR",
        "LA VICTORIA": "LIMA SUR",
        "MIRAFLORES": "LIMA SUR",
        "SAN BORJA": "LIMA SUR",
        "SANTIAGO DE SURCO": "LIMA SUR",
        "SURQUILLO": "LIMA SUR",
        "ATE VITARTE": "LIMA SUR",
        "ATE": "LIMA SUR",
        "CHACLACAYO": "LIMA SUR",
        "CIENEGUILLA": "LIMA SUR",
        "LA MOLINA": "LIMA SUR",
        "LURIGANCHO": "LIMA SUR",
        "SAN LUIS": "LIMA SUR",
        "CHORRILLOS": "LIMA SUR",
        "LURIN": "LIMA SUR",
        "PACHACAMAC": "LIMA SUR",
        "PUCUSANA": "LIMA SUR",
        "PUNTA HERMOSA": "LIMA SUR",
        "PUNTA NEGRA": "LIMA SUR",
        "SAN BARTOLO": "LIMA SUR",
        "SAN JUAN DE MIRAFLORES": "LIMA SUR",
        "SANTA MARIA DEL MAR": "LIMA SUR",
        "VILLA EL SALVADOR": "LIMA SUR",
        "VILLA MARIA DEL TRIUNFO": "LIMA SUR"
    }
    fechas_limite = {
        "ICA": "31/3/2025",
        "LIMA NORTE": "31/3/2025",
        "LIMA SUR": "31/3/2025",
        "AREQUIPA": "30/5/2025",
        "MOQUEGUA": "30/5/2025",
        "TACNA": "30/5/2025",
        "APURIMAC": "31/10/2025",
        "CUSCO": "31/10/2025",
        "MADRE DE DIOS": "31/10/2025",
        "PUNO": "31/10/2025",
        "AYACUCHO": "27/2/2026",
        "HUANCAVELICA": "27/2/2026",
        "HUANUCO": "27/2/2026",
        "JUNIN": "27/2/2026",
        "PASCO": "27/2/2026",
        "UCAYALI": "27/2/2026",
        "AMAZONAS": "30/6/2026",
        "ANCASH": "30/6/2026",
        "CAJAMARCA": "30/6/2026",
        "LA LIBERTAD": "30/6/2026",
        "SAN MARTIN": "30/6/2026",
        "LAMBAYEQUE": "31/8/2026",
        "LORETO": "31/8/2026",
        "PIURA": "31/8/2026",
        "TUMBES": "31/8/2026"
    }
    months = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"
    ]
    group = ["LIMA E ICA", "AREQUIPA, TACNA Y MOQUEGUA", "CUSCO, PUNO, MADRE DE DIOS Y APURIMAC", "JUNIN, AYACUCHO, HUANCAVELICA, HUANUCO, PASCO Y UCAYALI", "LA LIBERTAD, ANCASH, CAJAMARCA, SAN MARTIN, AMAZONAS Y LORETO", "LAMBAYEQUE, PIURA Y TUMBES", "LORETO"]
    teams = ["Lima e Ica", "Arequipa, Tacna y Moquegua", "Cusco, Puno, Madre de Dios y Apurimac", "Junín, Ayacucho, Huancavelica, Huánuco, Pasco y Ucayali", "La Libertad, Ancash, Cajamarca, San Martín, Amazonas y Loreto", "Lambayeque, Piura y Tumbes", "Loreto"]
    return (
        departamentos_region,
        distritos_region,
        fechas_limite,
        group,
        months,
        provincias_region,
        teams,
    )


@app.cell(hide_code=True)
def _(
    BytesIO,
    Cm,
    DocxTemplate,
    InlineImage,
    alt,
    datetime,
    departamentos_region,
    distritos_region,
    fechas_limite,
    io,
    mo,
    months,
    pd,
    provincias_region,
    re,
    requests,
    zipfile,
):
    class Texto:
        def __init__(self, text, size, align, kind):
            self.text = text
            self.size = str(size)
            self.align = align
            self.kind = kind
        def create(self):
            return mo.md(f"""<div style = "text-align: {self.align}; font-size: {self.size}px; font-weight: {self.kind};">{self.text}</div>""")

    def regiones_outofdate(fechas):
    # Esta funcion devuelve aquellas regiones para las que su fecha limite ya caduco.
        hoy = datetime.now()
        regiones_vencidas = [
            region for region, fecha_str in fechas.items()
            if datetime.strptime(fecha_str, "%d/%m/%Y") < hoy
        ]
        return regiones_vencidas

    def cleaner(data, codigo, registro):
    # Esta funcion elimina las filas donde el codigo y el registro se encuentren vacios, y donde ambos valores, se repitan. En caso haya filas con codigos repetidos y registros diferentes, se eliminan las filas con registros antiguos, quedandose unicamente con el que contenga el registro mas reciente. 
        data = data.dropna(subset = [codigo, registro], ignore_index = True).drop_duplicates(subset = [codigo, registro], ignore_index = True)
        repeated = data[codigo].duplicated(keep = False)
        no_repeats = data[~repeated].copy()
        repeats = data[repeated].copy()
        repeats["FECHA"] = pd.to_datetime(repeats[registro].str.split("-").str[2], format = '%d%m%y', errors = "coerce")
        repeats = repeats.sort_values(by = ["FECHA"], na_position = "first").drop_duplicates(codigo, keep = "last").drop("FECHA", axis = 1)
        return pd.concat([no_repeats, repeats], ignore_index = True)

    def obtener_region(departamento, provincia, distrito):
        if departamento and provincia and distrito:
            if departamento in ["PROV. CONST. DEL CALLAO", "LIMA", "CALLAO"]:
                region = distritos_region.get(distrito)
                if region:
                    return region
                region = provincias_region.get(provincia)
                if region:
                    return region
                return departamento
            else:
                return departamentos_region.get(departamento, departamento)
        else:
            return "REVISAR"

    def dateofmade(month, year):
    # Esta funcion convierte un mes y un año, en una fecha en formato 01/mes/año, teniendo en cuenta que para años menores a 1900, se obtara por establecer dicho año como año de fabricacion, es decir, el formato sera 01/mes/1900.
        return datetime(max(int(year), 1900), months.index(month.upper().strip()) + 1, 1).date()

    def replicate(series, data, tanks_data, tests_data):
    # Se extrae la informacion de las bases de datos referidos a registro de componentes y pruebas de hermeticidad, considerando para el ello, el numero de registro, y en caso no se encuentre coincidencias con este, se efectua una busqueda por el codigo, como segunda opcion.
        record = series.loc["REGISTRO"]
        code = int(series.loc["CODIGO OSINERGMIN"])
        data_tanks = tanks_data[tanks_data["NUMERO_REGISTRO_HIDROCARBURO"] == record]
        if data_tanks.empty:
            data_tanks = tanks_data[tanks_data["CODIGO_OSINERGMIN"] == code]
        data_tests = tests_data[tests_data["REGISTRO DE HIDROCARBURO"] == record].copy()
        if data_tests.empty:
            data_tests = tests_data[tests_data["CÓDIGO OSINERGMIN"] == code].copy()
        if not data_tests.empty:
            data_tests["FECHA FIN"] = pd.to_datetime(data_tests["FECHA FIN"], format='%d-%m-%Y', errors = "coerce").dt.date
    # Se extrae y almacena la informacion que no varia a nivel de codigo
        code_fixed = series.loc["CODIGO OSINERGMIN":"DIRECCION OPERATIVA"].tolist() + [series.loc["REGION"], series.loc["ACTIVIDAD"]]
        total_capacity = [series.loc["CAP.TOTAL CL (gln)"]]
        tanks = series.loc["TANQUE 1":"TANQUE 15"].values
        for j, i in enumerate(tanks, start = 1):
            if pd.notna(i):
    # Se extrae y almacena la informacion que no varia a nivel de tanque (registro de componentes)
                if not data_tanks.empty:
                    data_tank = data_tanks[data_tanks["NUMERO_TANQUE"] == j]
                    if not data_tank.empty:
                        month, year, location, material = data_tank[["MES_FABRICACION", "ANIO_FABRICACION", "TIPO_INSTALACION", "MATERIAL"]].fillna("").values[0]
                        date = dateofmade(month, year) if month and year else ""
                    else:
                        location = material = date = ""
                    tank_fixed = [material.upper(), location.upper(), date]
                else:
                    tank_fixed = [""] * 3
    # Se extrae y almacena la informacion que no varia a nivel de tanque (pruebas de hermeticidad)
                if not data_tests.empty:
                    data_tank = data_tests[data_tests["TANQUE NRO"] == j].copy()
                else:
                    data_tank = pd.DataFrame()
                slots = re.findall(r"(C\d+:.*?)(?=\sC\d+|$)", i.replace("\xa0", "")) or [i]
                for k in slots:
                    content = k.strip().split(":")
                    if len(content) == 3:
                        slot = int(content[0].strip("C"))
                        data.append(code_fixed + tank_fixed + total_capacity + [j, slot, content[1], content[2].upper()])
                    elif len(content) == 2:
                        slot = 1
                        data.append(code_fixed + tank_fixed + total_capacity + [j, slot, content[0], content[1].upper()])
                    else:
                        slot = ""
                        data.append(code_fixed + tank_fixed + total_capacity + [j, slot] + [""] * 9)
    # Se filtra la informacion de pruebas de hermeticidad a nivel de compartimiento, ya que un compartimiento puede haberse sometido a distintas pruebas, y es necesario solo trabajar con el resultado registrado mas reciente.
                    if slot and not data_tank.empty:
                        data_slot = data_tank[data_tank["NRO COMPARTIMIENTO"] == slot].sort_values(by = ["FECHA FIN"], na_position = "first").drop_duplicates(subset = ["NRO COMPARTIMIENTO"], keep = "last")
    # Se extrae y almacena la informacion que no varia a nivel de compartimiento (pruebas de hermeticidad)
                        if not data_slot.empty:
                            requirement_status, test_date, slot_result, pipes_result, final_result, certificate, inspector = data_slot[["ESTADO SOLICITUD", "FECHA FIN", "RESULTADO DE COMPARTIMIENTO", "RESULTADO DE TUBERÍA", "RESULTADO", "NRO DE CERTIFICADO", "ORGANISMO DE INSPECCIÓN"]].fillna("").values[0]
                        else:
                            requirement_status = test_date = slot_result = pipes_result = final_result = certificate = inspector = ""
                        data[-1] = data[-1] + [requirement_status, test_date, slot_result, pipes_result, final_result, certificate, inspector]
        return data

    def validate_content(product):
        if "GLP" in product or "PETROLEO" in product:
            return "GLP"
        elif "PRODUCTO" in product or product.strip() == "":
            return "S/P"
        else:
            return "CL"

    def tank_state(fabricacion, contenido):
    # Las modificaciones corresponden al hecho de que, casi en toda su totalidad, los grifos cuentan con componentes enterrados, por lo que, practicamente todos ellos deberian hacer la prueba de hermeticidad correspondiente. (21/05/2025)
        if contenido == "GLP":
            return "NO APLICA"
        else:
            if fabricacion:
                if fabricacion <= datetime(2022, 1, 12).date():
                    return "PRE EXISTENTE"
                else:
                    return "NUEVO"
            else:
                return ""

    def fecha_limite(estado, provincia, region):
        if estado not in ("NO APLICA", "NUEVO"):
            if provincia == "ALTO AMAZONAS":
                return datetime(2026, 6, 30).date()
            else:
                return datetime.strptime(fechas_limite.get(region), "%d/%m/%Y").date()
        else:
            return estado

    def tank_age(estado, fabricacion):
        if estado != "NO APLICA":
            if fabricacion:
                return round((datetime.today().date() - fabricacion).days / 365, 1)
            else:
                return ""
        else:
            return estado

    def registro(data):
    # Ya no hay cuarto estado de NO CORRESPONDE debido a que no se utilizara el filtro de ubicacion
        data_agent = data[
            (data["ESTADO DEL TANQUE"] != "NO APLICA")
        ]
        tanks = data_agent["TANQUE"].unique()
        key = []
        for i in tanks:
            material, ubicacion, fabricacion = data_agent[data_agent["TANQUE"] == i][["MATERIAL DEL TANQUE","UBICACIÓN DEL TANQUE", "FECHA DE FABRICACION"]].values[0]
            if material and ubicacion and fabricacion:
                key.append(2)
            elif material == "" and ubicacion == "" and fabricacion == "":
                key.append(0)
            else:
                key.append(1)
                break
        key = sum(key)
        if key == 2 * len(tanks):
            return 2
        elif key == 0:
            return 0
        else:
            return 1

    def hermeticidad(data):
    # Ya no hay cuarto estado de NO CORRESPONDE debido a que no se utilizara el filtro de ubicacion
        datest_agent = data[
            ((data["ESTADO DEL TANQUE"] == "PRE EXISTENTE") |
             (data["ESTADO DEL TANQUE"] == ""))
        ]
        result = (
            (datest_agent['RESULTADO DE COMPARTIMIENTO'] == 'SIN FUGA').all() and
            (datest_agent['RESULTADO DE TUBERÍA'] == 'SIN FUGA').all()
        )
        if result:
            return 2
        else:
            restdos = datest_agent[["RESULTADO DE COMPARTIMIENTO", "RESULTADO DE TUBERÍA"]].values.tolist()
            results = [elemento for sublista in restdos for elemento in sublista]
            if "SIN FUGA" not in results:
                return 0
            else:
                return 1

    @mo.cache
    def registro_hermeticidad(data):
        datos = []
        codes = data["CODIGO OSINERGMIN"].unique()
        for i in codes:
            code_data = data[data["CODIGO OSINERGMIN"] == i]
            datos.append([i, code_data["OFICINA REGIONAL"].iloc[0], registro(code_data), hermeticidad(code_data)])
        return pd.DataFrame(datos, columns = ["CODIGO OSINERGMIN", "OFICINA REGIONAL", "REGISTRO DE INFORMACION", "PRUEBAS DE HERMETICIDAD"])

    @mo.cache
    def load_tanksdata(filename, sheetname):
        return pd.read_excel(filename, sheet_name = sheetname)

    @mo.cache
    def load_testsdata(filename, sheetname):
        data = pd.read_excel(filename, sheet_name = sheetname)
        data = data[
            (data["ESTADO SOLICITUD"] == "PRUEBA CONCLUIDA") |
            (data["ESTADO SOLICITUD"] == "EN REGISTRO")
        ]
    # Filtro para extraer informacion correspondiente a las actividades a analizar
        validate = ["050", "056", "106", "107", "EESS", "GRIF", "EGNV", "EMIX", "ICA"]
        regex = "|".join(validate)
        data["VALIDACION"] = data["REGISTRO DE HIDROCARBURO"].str.contains(regex).astype(int)
        data = data[data["VALIDACION"] == 1].drop("VALIDACION", axis = 1)
    # Filtro para retirar informacion correspondiente a las actividades que no se desea analizar
        data["VALIDACION"] = data["REGISTRO DE HIDROCARBURO"].str.split("-").str[1]
        data = data[(data["VALIDACION"] != "CDFJ") & (data["VALIDACION"] != "051")].drop("VALIDACION", axis = 1)
        return data

    @mo.cache
    def load_urlsdata(code_title, record_title):
        urls = ["https://pvo.osinergmin.gob.pe/msfh5/registroHidrocarburos.xhtml?method=excel&actividad=" + i for i in ["1", "2", "5", "6"]]
        columns = ["CODIGO OSINERGMIN", "REGISTRO", "RUC", "RAZON SOCIAL", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "DIRECCION OPERATIVA"] + ["TANQUE " + str(i) for i in range(1, 16)] + ["CAP.TOTAL CL (gln)", "ACTIVIDAD"]
        activities = ["050 - ESTACION DE SERVICIOS / GRIFOS", "056 - ESTACION DE SERVICIO CON GASOCENTRO DE GLP", "106 - ESTACION DE SERVICIOS CON ESTABLECIMIENTO DE VENTA DE GNV", "107 - ESTACION DE SERVICIO CON GASOCENTRO DE GLP Y ESTABLECIMIENTO DE VENTA DE GNV"]
        data1 = []
        for i in range(4):
            response = requests.get(urls[i])
            data0 = pd.read_html(BytesIO(response.content))[0]
            data0["ACTIVIDAD"] = activities[i]
            data1.append(data0.reindex(columns = columns))
        data1 = pd.concat(data1, ignore_index = True)
        data1["CODIGO OSINERGMIN"] = data1["CODIGO OSINERGMIN"].str[2:-1]
        data1 = cleaner(data1, code_title, record_title)
        data1["REGION"] = data1.apply(lambda row: obtener_region(row["DEPARTAMENTO"], row["PROVINCIA"], row["DISTRITO"]), axis = 1)
        return data1

    @mo.cache
    def load_matrix(data_url, data_tank, data_test):
        data = []
        for _, _i in data_url.iterrows():
            replicate(_i, data, data_tank, data_test)
        data = pd.DataFrame(data, columns = ["CODIGO OSINERGMIN", "REGISTRO", "RUC", "RAZON SOCIAL", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "DIRECCION OPERATIVA", "OFICINA REGIONAL", "ACTIVIDAD", "MATERIAL DEL TANQUE", "UBICACIÓN DEL TANQUE", "FECHA DE FABRICACION", "CAPACIDAD TOTAL DE CL (GLN)", "TANQUE", "COMPARTIMIENTO", "CAPACIDAD (GLN)", "PRODUCTO", "ESTADO DE SOLICITUD", "FECHA DE PRUEBA", "RESULTADO DE COMPARTIMIENTO", "RESULTADO DE TUBERÍA", "RESULTADO FINAL", "CERTIFICADO", "ORGANISMO DE INSPECCION"])
        data["CODIGO OSINERGMIN"] = data["CODIGO OSINERGMIN"].astype(int)
        data["CAPACIDAD (GLN)"] = data["CAPACIDAD (GLN)"].astype(float)
        data = pd.concat([data, notinlists(data, data_test, data_tank)], ignore_index = True)
        data["CONTENIDO"] = data["PRODUCTO"].apply(lambda x: validate_content(x))
        data["ESTADO DEL TANQUE"] = data.apply(lambda row: tank_state(row["FECHA DE FABRICACION"], row["CONTENIDO"]), axis = 1)
        data["FECHA LIMITE"] = data.apply(lambda row: fecha_limite(row["ESTADO DEL TANQUE"], row["PROVINCIA"], row["OFICINA REGIONAL"]), axis = 1)
        data["EDAD DEL TANQUE"] = data.apply(lambda row: tank_age(row["ESTADO DEL TANQUE"], row["FECHA DE FABRICACION"]), axis = 1)
        return data

    def plot_bars(data, state1, state2, state3):
        large = data.melt(
            id_vars = "OFICINA REGIONAL",
            value_vars = [
                state1, 
                state2, 
                state3,
            ],
            var_name = "ESTADO", 
            value_name = "CANTIDAD"
        )
        orden_regiones = (
            large.groupby("OFICINA REGIONAL")["CANTIDAD"]
            .sum()
            .sort_values(ascending = False)
            .index
            .tolist()
        )
        colores_personalizados = {
            state1: "#0000FF",   
            state2: "#ff1100",  
            state3: "#ffffff"         
        }
        graph = alt.Chart(large).mark_bar(opacity = 0.7, stroke = "black", strokeWidth = 0.5).encode(
            x = alt.X("OFICINA REGIONAL:N", title = "OFICINA REGIONAL", sort = orden_regiones),
            y = alt.Y("CANTIDAD:Q", title = "CANTIDAD DE AGENTES OPERATIVOS", stack = "zero"),
            color = alt.Color(
                "ESTADO:N",
                title = "ESTADO",
                scale = alt.Scale(
                    domain = list(colores_personalizados.keys()), 
                    range = list(colores_personalizados.values())
                ),
                legend = alt.Legend(
                    orient = "top-right",
                    padding = 0,
                    labelLimit = 500,
                )
            )
        ).properties(
            title = "ESTADO DEL CUMPLIMIENTO DEL D.S. 001-2022-MINEM-EM POR OFICINA REGIONAL"
        )
        return mo.ui.altair_chart(graph)

    def notinlists(matrix, data_tests, data_tanks):
        columns = ["CODIGO OSINERGMIN", "REGISTRO", "RUC", "RAZON SOCIAL", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "DIRECCION OPERATIVA", "OFICINA REGIONAL", "ACTIVIDAD", "MATERIAL DEL TANQUE", "UBICACIÓN DEL TANQUE", "FECHA DE FABRICACION", "CAPACIDAD TOTAL DE CL (GLN)", "TANQUE", "COMPARTIMIENTO", "CAPACIDAD (GLN)", "PRODUCTO", "ESTADO DE SOLICITUD", "FECHA DE PRUEBA", "RESULTADO DE COMPARTIMIENTO", "RESULTADO DE TUBERÍA", "RESULTADO FINAL", "CERTIFICADO", "ORGANISMO DE INSPECCION"]
        codes = list(set(data_tests["CÓDIGO OSINERGMIN"]) - set(matrix["CODIGO OSINERGMIN"]))
        data = []
        for code in codes:
            df = data_tests[data_tests["CÓDIGO OSINERGMIN"] == code].copy()
            df["FECHA FIN"] = pd.to_datetime(df["FECHA FIN"], format='%d-%m-%Y', errors = "coerce").dt.date
            fila = df.iloc[0]
            record = fila["REGISTRO DE HIDROCARBURO"]
            departamento, provincia, distrito = fila["DEPARTAMENTO":"DISTRITO"].tolist()
            code_fixed = [fila["CÓDIGO OSINERGMIN"], record, "", fila["RAZÓN SOCIAL"], departamento, provincia, distrito, fila["DIRECCIÓN OPERATIVA"], obtener_region(departamento, provincia, distrito), activity(record)]
            fd = data_tanks[data_tanks["CODIGO_OSINERGMIN"] == code]
            available_tanks = sorted(df["TANQUE NRO"].unique().tolist())
            dato = []
            for i in available_tanks:
                df1 = df[df["TANQUE NRO"] == i]
                if not fd.empty:
                    data_tank = fd[fd["NUMERO_TANQUE"] == i]
                    if not data_tank.empty:
                        month, year, location, material = data_tank[["MES_FABRICACION", "ANIO_FABRICACION", "TIPO_INSTALACION", "MATERIAL"]].fillna("").values[0]
                        date = dateofmade(month, year) if month and year else ""
                    else:
                        location = material = date = ""
                    tank_fixed = [material.upper(), location.upper(), date]
                else:
                    tank_fixed = [""] * 3
                available_slots = sorted(df1["NRO COMPARTIMIENTO"].unique().tolist())
                for j in available_slots:
                    df2 = df1[df1["NRO COMPARTIMIENTO"] == j].sort_values(by = ["FECHA FIN"], na_position = "first").drop_duplicates(subset = ["NRO COMPARTIMIENTO"], keep = "last")
                    row = df2.iloc[0]
                    dato.append(code_fixed + tank_fixed + ["", i, j, row["CAPACIDAD DE COMPARTIMIENTO"], row["PRODUCTO"], row["ESTADO SOLICITUD"], row["FECHA FIN"], row["RESULTADO DE COMPARTIMIENTO"], row["RESULTADO DE TUBERÍA"], row["RESULTADO"], row["NRO DE CERTIFICADO"], row["ORGANISMO DE INSPECCIÓN"]])
            dato = pd.DataFrame(dato, columns = columns)
            dato["CAPACIDAD TOTAL DE CL (GLN)"] = dato["CAPACIDAD (GLN)"].sum()
            data.append(dato)
        return pd.concat(data)

    def activity(record):
        fragments = record.split("-")
        if len(fragments) == 3:
            code = fragments[1]
            if code == "050":
                return "050 - ESTACION DE SERVICIOS / GRIFOS"
            elif code == "056":
                return "056 - ESTACION DE SERVICIO CON GASOCENTRO DE GLP"
            elif code == "106":
                return "106 - ESTACION DE SERVICIOS CON ESTABLECIMIENTO DE VENTA DE GNV"
            elif code == "107":
                return "107 - ESTACION DE SERVICIO CON GASOCENTRO DE GLP Y ESTABLECIMIENTO DE VENTA DE GNV"
            else:
                return "REVISAR"
        elif "EGNV" in fragments:
            return "106 - ESTACION DE SERVICIOS CON ESTABLECIMIENTO DE VENTA DE GNV"
        elif "EMIX" in fragments or "ICA" in fragments or "EESS" in fragments or "GRIF" in fragments:
            return "050 - ESTACION DE SERVICIOS / GRIFOS"
        else:
            return "REVISAR"

    def generatedocs(region, lugar, regiones, mes, anio, cronograma, datos_or, telefono, jor, cargo, result, matrix):
        with mo.status.spinner(title = 'Filtrando agentes', subtitle = 'espere ...', remove_on_exit = True) as _spinner:
            _codes = result[(result["OFICINA REGIONAL"] == region.value) & ((result["PRUEBAS DE HERMETICIDAD"] == 0) & (result["REGISTRO DE INFORMACION"] == 0))]["CODIGO OSINERGMIN"]
            _data_filter = matrix[matrix["CODIGO OSINERGMIN"].isin(_codes)].reset_index(drop = True)
            _data_filter = _data_filter[["CODIGO OSINERGMIN", "REGISTRO", "RUC", "RAZON SOCIAL", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "DIRECCION OPERATIVA", "OFICINA REGIONAL", "ACTIVIDAD"]].drop_duplicates(subset = ["CODIGO OSINERGMIN"], ignore_index = True)

        buffer_zip = io.BytesIO()
        with zipfile.ZipFile(buffer_zip, mode = "w", compression = zipfile.ZIP_DEFLATED) as zipf:

            with mo.status.progress_bar(total = _data_filter.shape[0], title = "Generando oficio", completion_title = "Oficios generados", completion_subtitle = str(_data_filter.shape[0]) + " en total", remove_on_exit = True) as bar:
                for _, _row in _data_filter.iterrows():
                    razon_social = _row["RAZON SOCIAL"]
                    direccion = _row["DIRECCION OPERATIVA"]
                    registro = _row["REGISTRO"]
                    distrito = _row["DISTRITO"]
                    provincia = _row["PROVINCIA"]
                    departamento = _row["DEPARTAMENTO"]
                    codigo = _row["CODIGO OSINERGMIN"]

                    _document = DocxTemplate("PLANTILLA.docx")
                    _context = {
                        "LUGAR": lugar.value,
                        "RAZON_SOCIAL": razon_social,
                        "DIRECCION": direccion,
                        "REGIONES": regiones,
                        "MES": mes,
                        "ANIO": anio,
                        "CRONOGRAMA": InlineImage(_document, cronograma, width=Cm(15.5)),
                        "REGISTRO": registro,
                        "DISTRITO": distrito,
                        "PROVINCIA": provincia,
                        "DEPARTAMENTO": departamento,
                        "DATOS_OR": datos_or,
                        "TELEFONO": telefono.value,
                        "JOR": jor.value,
                        "CARGO": cargo.value,
                    }
                    _document.render(_context)               
                    file_stream = io.BytesIO()
                    _document.save(file_stream)
                    file_stream.seek(0)
                    zipf.writestr(f"{codigo}.docx", file_stream.read())

                    bar.update(subtitle = "codigo: " + str(codigo))

        with mo.status.spinner(title = 'Generando ZIP', subtitle = 'espere ...', remove_on_exit = True) as _spinner:
            buffer_zip.seek(0)
            zip_download = mo.download(
                data = buffer_zip.read(),
                filename = "OFICIOS_" + region.value + ".zip",
                mimetype = "application/zip",
                label = "DESCARGAR OFICIOS A EMITIR"
            )
            return mo.output.replace(zip_download)
    return (
        Texto,
        generatedocs,
        load_matrix,
        load_tanksdata,
        load_testsdata,
        load_urlsdata,
        plot_bars,
        regiones_outofdate,
        registro_hermeticidad,
    )


@app.cell
def _(Texto, mo):
    _texto = "HERMETICIDAD DE TANQUES DE CL y OPDH A NIVEL NACIONAL"
    mo.output.append(Texto(_texto, 36, "center", "bold").create())
    _texto = "** aplica solo para tanques ubicados en estaciones de servicios **"
    mo.output.append(Texto(_texto, 16, "center", "bold").create())
    mo.output.append(mo.md('---'))
    mo.output.append(mo.md("<br>"))
    _texto = "Por medio del Decreto Supremo 001-2022-MINEM-EM se aprobo la norma para la inspección periódica de hermeticidad de tuberías y tanques enterrados que almacenan combustibles líquidos y otros productos derivados de los hidrocarburos, la cual establece las condiciones técnicas mínimas para la Inspección periódica de Hermeticidad de los tanques enterrados y las tuberías enterradas que almacenan Combustibles Líquidos y Otros Productos Derivados de los Hidrocarburos. Así mismo, dicho decreto establecio que la inspección descrita debera de efectuarse bajo el siguiente cronograma:"
    mo.output.append(Texto(_texto, 16, "justify", "normal").create())
    mo.output.append(mo.image(
        src = "CRONOGRAMA.png",
        width = "100%",
    ))
    mo.output.append(mo.md("<br>"))
    _texto = "MATRIZ GENERAL DE DATOS ACTUALIZADA"
    download_2 = mo.ui.switch(label = "Exportar matriz general de datos")
    mo.output.append(mo.hstack([Texto(_texto, 26, "left", "bold").create(), download_2], justify = "space-between", align = "stretch"))
    return (download_2,)


@app.cell
def _(
    load_matrix,
    load_tanksdata,
    load_testsdata,
    load_urlsdata,
    mo,
    registro_hermeticidad,
):
    with mo.status.spinner(title = 'Cargando datos', subtitle = 'de tanques de almacenamiento de CL y OPDH', remove_on_exit = True) as _spinner:
        _data_tanks = load_tanksdata("DATA TANQUES.xlsx", "Exportar Hoja de Trabajo")
    with mo.status.spinner(title = 'Cargando datos', subtitle = 'de pruebas de hermeticidad ejecutadas a nivel nacional', remove_on_exit = True) as _spinner:
        _data_tests = load_testsdata("DATA PRUEBAS.xlsx", "PH - PRUEBAS HERMETICIDAD")
    with mo.status.spinner(title = 'Cargando datos', subtitle = 'de agentes habilitados a nivel nacional', remove_on_exit = True) as _spinner:
        _data_urls = load_urlsdata("CODIGO OSINERGMIN", "REGISTRO")
    with mo.status.spinner(title = 'Consolidando datos', subtitle = 'y generando matriz de datos', remove_on_exit = True) as _spinner:
        matriz = load_matrix(_data_urls, _data_tanks, _data_tests)
    with mo.status.spinner(title = 'Filtrando agentes que acreditaron', subtitle = 'registro de informacion y hermeticidad', remove_on_exit = True) as _spinner:
        result = registro_hermeticidad(matriz)
    mo.output.replace(matriz)
    return matriz, result


@app.cell
def _(download_2, io, matriz, mo):
    if download_2.value:
        with mo.status.spinner(title = 'Generando Excel', subtitle = 'espere ...', remove_on_exit = True) as _spinner:
            _buffer = io.BytesIO()
            matriz.to_excel(_buffer, index = False, engine = "openpyxl")
            _buffer.seek(0)
            _excel_download = mo.download(
                data = _buffer.read(),
                filename = "MATRIZ.xlsx",
                mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                label = "Descargar Excel"
            )
            mo.output.append(_excel_download)
    return


@app.cell
def _(Texto, mo, pd, plot_bars, result):
    mo.output.append(mo.md("<br>"))
    _texto = "RCD N° 010-2023-OS/CD"
    mo.output.append(Texto(_texto, 26, "left", "bold").create())
    _texto = "Establece las disposiciones a las que se sujeta el uso obligatorio de la plataforma denominada “Soluciones Tecnológicas de Gestión de Componentes de Almacenamiento para el registro del proceso de pruebas de hermeticidad en sistemas de tanques enterrados” (en adelante, la Plataforma), de conformidad con la “Norma para la Inspección Periódica de Hermeticidad de tuberías y tanques enterrados que almacenan Combustibles Líquidos y Otros Productos Derivados de los Hidrocarburos”, aprobada por Decreto Supremo N° 001-2022-EM."
    mo.output.append(Texto(_texto, 16, "justify", "normal").create())

    _regiones = result["OFICINA REGIONAL"].unique()
    _data = []
    for _i in _regiones:
        _data_region = result[result["OFICINA REGIONAL"] == _i]
        _data.append([_i, _data_region.shape[0], _data_region[_data_region["REGISTRO DE INFORMACION"] == 2].shape[0], _data_region[_data_region["REGISTRO DE INFORMACION"] == 1].shape[0], _data_region[_data_region["REGISTRO DE INFORMACION"] == 0].shape[0]])
    _data_registro = pd.DataFrame(_data, columns = ["OFICINA REGIONAL", "AGENTES OPERATIVOS", "REGISTRO DE INFORMACION COMPLETA", "REGISTRO DE INFORMACION INCOMPLETA", "REGISTRO DE INFORMACION NULA"]).sort_values(by = "OFICINA REGIONAL", ascending = True, ignore_index = True)
    mo.output.append(_data_registro)
    mo.output.append(plot_bars(_data_registro, "REGISTRO DE INFORMACION COMPLETA", "REGISTRO DE INFORMACION INCOMPLETA", "REGISTRO DE INFORMACION NULA"))
    return


@app.cell
def _(Texto, mo, pd, plot_bars, result):
    mo.output.append(mo.md("<br>"))
    _texto = "RCD N° 151-2024-OS/CD"
    mo.output.append(Texto(_texto, 26, "left", "bold").create())
    _texto = "Establece el procedimiento técnico para la implementación de la “Norma para la Inspección Periódica de Hermeticidad de tuberías y tanques enterrados que almacenan Combustibles Líquidos y Otros Productos Derivados de los Hidrocarburos”, aprobada mediante el artículo 10 del Decreto Supremo N° 001-2022-EM; que contiene las pautas a seguir para que las personas jurídicas interesadas puedan inscribirse en el Registro de Empresas Inspectoras que realizan las pruebas de Inspección de Hermeticidad del Sistema de Tanques Enterrados (STE), que las habilite a realizar inspecciones para verificar la hermeticidad de tanques y tuberías enterrados de los agentes fiscalizados que almacenan combustibles líquidos y otros productos derivados de los hidrocarburos."
    mo.output.append(Texto(_texto, 16, "justify", "normal").create())
    _regiones = result["OFICINA REGIONAL"].unique()
    _data = []
    for _i in _regiones:
        _data_region = result[result["OFICINA REGIONAL"] == _i]
        _data.append([_i, _data_region.shape[0], _data_region[_data_region["PRUEBAS DE HERMETICIDAD"] == 2].shape[0], _data_region[_data_region["PRUEBAS DE HERMETICIDAD"] == 1].shape[0], _data_region[_data_region["PRUEBAS DE HERMETICIDAD"] == 0].shape[0]])
    _data_hermeticidad = pd.DataFrame(_data, columns = ["OFICINA REGIONAL", "AGENTES OPERATIVOS", "HERMETICIDAD COMPLETA", "HERMETICIDAD INCOMPLETA", "HERMETICIDAD NULA"]).sort_values(by = "OFICINA REGIONAL", ascending = True, ignore_index = True)
    mo.output.append(_data_hermeticidad)
    mo.output.append(plot_bars(_data_hermeticidad, "HERMETICIDAD COMPLETA", "HERMETICIDAD INCOMPLETA", "HERMETICIDAD NULA"))
    return


@app.cell
def _(Texto, alt, intervalos, matriz, mo, pd):
    mo.output.append(mo.md("<br>"))
    _texto = "MAPAS DE CALOR PARA TANQUES HERMETICOS Y NO HERMETICOS A NIVEL NACIONAL"
    mo.output.append(Texto(_texto, 26, "left", "bold").create())

    _data = matriz[[
        "CODIGO OSINERGMIN", 
        "TANQUE",
        "CAPACIDAD (GLN)",
        "RESULTADO DE COMPARTIMIENTO", 
        "EDAD DEL TANQUE"
    ]].copy().reset_index(drop = True)
    _data = _data[
        _data["RESULTADO DE COMPARTIMIENTO"].notna() & 
        (_data["RESULTADO DE COMPARTIMIENTO"].str.strip() != "") &
        _data["EDAD DEL TANQUE"].notna() & 
        (_data["EDAD DEL TANQUE"].str.strip() != "") &
        (_data["EDAD DEL TANQUE"] != "NO APLICA")
    ].reset_index(drop = True)
    _datos = []
    _codes = _data["CODIGO OSINERGMIN"].unique()
    for _i in _codes:
        _data_code = _data[_data["CODIGO OSINERGMIN"] == _i]
        _tanks = _data_code["TANQUE"].unique()
        for _j in _tanks:
            _data_tank = _data_code[_data_code["TANQUE"] == _j]
            if (_data_tank['RESULTADO DE COMPARTIMIENTO'] == 'SIN FUGA').all():
                _result = "HERMETICO"
            else:
                _result = "NO HERMETICO"
            _datos.append(list(_data_tank.values[0]) + [_result, _data_tank["CAPACIDAD (GLN)"].sum()])
    _datos = pd.DataFrame(_datos, columns = [
        "CODIGO OSINERGMIN",  
        "TANQUE",
        "CAPACIDAD (GLN)",
        "RESULTADO DE COMPARTIMIENTO", 
        "EDAD DEL TANQUE",
        "RESULTADO",
        "CAPACIDAD DEL TANQUE (GLN)"
    ]).drop(columns = ["CODIGO OSINERGMIN", "TANQUE", "CAPACIDAD (GLN)", "RESULTADO DE COMPARTIMIENTO"])

    _nohermeticos = _datos[(_datos["EDAD DEL TANQUE"] < 100) & (_datos["RESULTADO"] == "NO HERMETICO")].drop(columns = ["RESULTADO"])
    _nohermeticos["RANGO_EDAD"] = pd.cut(_nohermeticos["EDAD DEL TANQUE"], bins = intervalos(_nohermeticos["EDAD DEL TANQUE"].max()))
    _nohermeticos["RANGO_CAPACIDAD"] = pd.cut(_nohermeticos["CAPACIDAD DEL TANQUE (GLN)"], bins = intervalos(_nohermeticos["CAPACIDAD DEL TANQUE (GLN)"].max()))
    _nohermeticos = _nohermeticos.groupby(["RANGO_EDAD", "RANGO_CAPACIDAD"], observed = False).size().reset_index(name = "CANTIDAD")

    _orden_edad = _nohermeticos["RANGO_EDAD"].cat.categories.astype(str).tolist()
    _orden_capacidad = _nohermeticos["RANGO_CAPACIDAD"].cat.categories.astype(str).tolist()
    _nohermeticos["RANGO_EDAD"] = _nohermeticos["RANGO_EDAD"].astype(str)
    _nohermeticos["RANGO_CAPACIDAD"] = _nohermeticos["RANGO_CAPACIDAD"].astype(str)

    _heatmap_nh = alt.Chart(_nohermeticos).mark_rect().encode(
        x=alt.X("RANGO_EDAD:N", title="EDAD DEL TANQUE (AÑOS)", sort=_orden_edad),
        y=alt.Y("RANGO_CAPACIDAD:N", title="CAPACIDAD DEL TANQUE (GALONES)", sort=_orden_capacidad[::-1]),
        color=alt.Color("CANTIDAD:Q", scale=alt.Scale(scheme="reds"), title="CANTIDAD"),
        tooltip=["RANGO_EDAD", "RANGO_CAPACIDAD", "CANTIDAD"]
    ).properties(
        title="MAPA DE CALOR PARA TANQUES NO HERMETICOS"
    )

    _hermeticos = _datos[(_datos["EDAD DEL TANQUE"] < 100) & (_datos["RESULTADO"] == "HERMETICO")].drop(columns = ["RESULTADO"])
    _hermeticos["RANGO_EDAD"] = pd.cut(_hermeticos["EDAD DEL TANQUE"], bins = intervalos(_hermeticos["EDAD DEL TANQUE"].max()))
    _hermeticos["RANGO_CAPACIDAD"] = pd.cut(_hermeticos["CAPACIDAD DEL TANQUE (GLN)"], bins = intervalos(_hermeticos["CAPACIDAD DEL TANQUE (GLN)"].max()))
    _hermeticos = _hermeticos.groupby(["RANGO_EDAD", "RANGO_CAPACIDAD"], observed = False).size().reset_index(name = "CANTIDAD")

    _orden_edad = _hermeticos["RANGO_EDAD"].cat.categories.astype(str).tolist()
    _orden_capacidad = _hermeticos["RANGO_CAPACIDAD"].cat.categories.astype(str).tolist()
    _hermeticos["RANGO_EDAD"] = _hermeticos["RANGO_EDAD"].astype(str)
    _hermeticos["RANGO_CAPACIDAD"] = _hermeticos["RANGO_CAPACIDAD"].astype(str)

    _heatmap_sh = alt.Chart(_hermeticos).mark_rect().encode(
        x=alt.X("RANGO_EDAD:N", title="EDAD DEL TANQUE (AÑOS)", sort=_orden_edad),
        y=alt.Y("RANGO_CAPACIDAD:N", title="CAPACIDAD DEL TANQUE (GALONES)", sort=_orden_capacidad[::-1]),
        color=alt.Color("CANTIDAD:Q", scale=alt.Scale(scheme="blues"), title="CANTIDAD"),
        tooltip=["RANGO_EDAD", "RANGO_CAPACIDAD", "CANTIDAD"]
    ).properties(
        title="MAPA DE CALOR PARA TANQUES HERMETICOS"
    )
    mo.output.append(mo.hstack([mo.ui.altair_chart(_heatmap_sh), mo.ui.altair_chart(_heatmap_nh)], justify = "space-between", align = "center", widths = [1, 1]))
    return


@app.cell
def _(np):
    def intervalos(_number):
        _values = np.arange(0, np.ceil(_number), np.ceil(np.ceil(_number) / 11))
        if _values[-1] < _number:
            _values = np.concatenate((_values, np.array([_values[-1] + np.ceil(np.ceil(_number) / 11)])))
        return _values
    return (intervalos,)


@app.cell
def _(Texto, mo):
    mo.output.append(mo.md("<br>"))
    _texto = "ANALISIS REGIONAL"
    download_1 = mo.ui.switch(label = "Exportar analisis")
    outofdate = mo.ui.switch(value = True, label = "Analizar solo a los que estan fuera del plazo correspondiente")
    modificate = mo.ui.switch(label = "Para oficios")
    mo.output.append(mo.hstack([Texto(_texto, 26, "left", "bold").create(), outofdate, modificate, download_1], justify = "space-between", align = "stretch"))
    return download_1, modificate, outofdate


@app.cell
def _(fechas_limite, matriz, mo, outofdate, regiones_outofdate, result):
    region = mo.ui.dropdown(options = sorted(regiones_outofdate(fechas_limite)) if outofdate.value else sorted(result["OFICINA REGIONAL"].unique().tolist()),
                            value = "LIMA NORTE",
                            label = "OFICINA REGIONAL: ",
                            full_width = True)
    analisis = mo.ui.dropdown(options = ["HERMETICIDAD DE TANQUES", "REGISTRO DE INFORMACION", "HERMETICIDAD Y REGISTRO"],
                            value = "HERMETICIDAD DE TANQUES",
                            label = "ANALIZAR: ",
                            full_width = True)
    estado = mo.ui.dropdown(options = {"COMPLETA":2, "INCOMPLETA":1, "NULA":0},
                            value = "NULA",
                            label = "ESTADO: ",
                            full_width = True)
    matrix = matriz[matriz["OFICINA REGIONAL"].isin(regiones_outofdate(fechas_limite))].copy() if outofdate.value else matriz.copy()
    mo.output.append(mo.hstack([region, analisis, estado], widths = [1, 1, 1, 1], gap = 2, justify = "space-between", align = "center"))
    return analisis, estado, matrix, region


@app.cell
def _(
    analisis,
    download_1,
    estado,
    io,
    matrix,
    mo,
    modificate,
    region,
    result,
):
    if analisis.value == "HERMETICIDAD DE TANQUES":
        _codes = result[(result["OFICINA REGIONAL"] == region.value) & (result["PRUEBAS DE HERMETICIDAD"] == estado.value)]["CODIGO OSINERGMIN"]
    elif analisis.value == "REGISTRO DE INFORMACION":
        _codes = result[(result["OFICINA REGIONAL"] == region.value) & (result["REGISTRO DE INFORMACION"] == estado.value)]["CODIGO OSINERGMIN"]
    else:
        _codes = result[(result["OFICINA REGIONAL"] == region.value) & ((result["PRUEBAS DE HERMETICIDAD"] == estado.value) & (result["REGISTRO DE INFORMACION"] == estado.value))]["CODIGO OSINERGMIN"]
    _data_filter = matrix[matrix["CODIGO OSINERGMIN"].isin(_codes)].reset_index(drop = True)
    _data_filter = _data_filter[["CODIGO OSINERGMIN", "REGISTRO", "RUC", "RAZON SOCIAL", "DEPARTAMENTO", "PROVINCIA", "DISTRITO", "DIRECCION OPERATIVA", "OFICINA REGIONAL", "ACTIVIDAD"]].drop_duplicates(subset = ["CODIGO OSINERGMIN"], ignore_index = True) if modificate.value else _data_filter
    mo.output.append(_data_filter)

    if download_1.value:
        with mo.status.spinner(title = 'Generando Excel', subtitle = 'espere ...', remove_on_exit = True) as _spinner:
            _buffer = io.BytesIO()
            _data_filter.to_excel(_buffer, index = False, engine = "openpyxl")
            _buffer.seek(0)
            _excel_download = mo.download(
                data = _buffer.read(),
                filename = region.value + "_" + analisis.value + "_" + str(estado.value) + ".xlsx",
                mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                label = "Descargar Excel"
            )
            mo.output.append(_excel_download)
    return


@app.cell
def _(Texto, fechas_limite, group, mo, months, region, teams):
    mo.output.append(mo.md("<br>"))
    _texto = "OFICIOS REGIONALES"
    mo.output.append(Texto(_texto, 26, "left", "bold").create())
    _texto = "Edite los siguientes parametros para generar los oficios correspondientes a aquellos agentes que no han acreditado el registro de informacion y prueba de hermeticidad alguna, y que son competencia de la oficina regional previamente seleccionada:"
    mo.output.append(Texto(_texto, 16, "justify", "normal").create())

    lugar = mo.ui.text(label = "LUGAR: ", full_width = True, placeholder = "Miraflores")

    for _i in range(len(group)):
        if region.value.split(" ")[0] in group[_i]:
            regiones = teams[_i]
            break

    _, mes, anio = fechas_limite.get(region.value).split("/")
    mes = months[int(mes) - 1].lower()

    cronograma = "GRUPOS/G" + str(_i + 1) + ".png"

    datos_or = "Oficina Regional de " + region.value.title()

    telefono = mo.ui.text(label = "TELEFONO: ", full_width = True, placeholder = "123456789")

    jor = mo.ui.text(label = "FIRMANTE: ", full_width = True, placeholder = "Nombres y apellidos")

    cargo = mo.ui.text(label = "PUESTO: ", full_width = True, placeholder = "Jefe de oficina regional")

    mo.output.append(mo.hstack([lugar, telefono, jor, cargo], widths = [1, 1, 1, 1], gap = 2, justify = "space-between", align = "end"))
    return (
        anio,
        cargo,
        cronograma,
        datos_or,
        jor,
        lugar,
        mes,
        regiones,
        telefono,
    )


@app.cell
def _(
    anio,
    cargo,
    cronograma,
    datos_or,
    generatedocs,
    jor,
    lugar,
    matrix,
    mes,
    mo,
    region,
    regiones,
    result,
    telefono,
):
    button_1 = mo.ui.run_button(label = f"{mo.icon('eos-icons:rotating-gear', size = 15, color = "black")} GENERAR OFICIOS A EMITIR", full_width = True, on_change = lambda _: generatedocs(region, lugar, regiones, mes, anio, cronograma, datos_or, telefono, jor, cargo, result, matrix), kind = "danger")
    return (button_1,)


@app.cell
def _(button_1, mo):
    if button_1.value:
        mo.output.clear()
    else:
        mo.output.replace(button_1)
    return


@app.cell
def _(Texto, mo):
    mo.output.append(mo.md("<br>"))
    optional_1 = mo.ui.switch(label = "Todas las regiones")
    _texto = "ANALISIS CORRELACIONAL"
    mo.output.append(mo.hstack([Texto(_texto, 26, "left", "bold").create(), optional_1], justify = "space-between", align = "stretch"))
    _texto = "La siguiente grafica, representa la relacion entre la edad del tanque y su capacidad:"
    mo.output.append(Texto(_texto, 16, "justify", "normal").create())
    return (optional_1,)


@app.cell
def _(alt, matrix, matriz, mo, optional_1, outofdate, pd, region):
    if optional_1.value:
        _data_filter = matrix if outofdate.value else matriz.copy()
    else:
        _data_filter = matriz[matriz["OFICINA REGIONAL"] == region.value].copy()
    _regiones = str(sorted(_data_filter["OFICINA REGIONAL"].unique().tolist()))
    _data = _data_filter[[
        "CODIGO OSINERGMIN", 
        "TANQUE",
        "CAPACIDAD (GLN)",
        "RESULTADO DE COMPARTIMIENTO", 
        "EDAD DEL TANQUE"
    ]].copy().reset_index(drop = True)
    _data = _data[
        _data["RESULTADO DE COMPARTIMIENTO"].notna() & 
        (_data["RESULTADO DE COMPARTIMIENTO"].str.strip() != "") &
        _data["EDAD DEL TANQUE"].notna() & 
        (_data["EDAD DEL TANQUE"].str.strip() != "") &
        (_data["EDAD DEL TANQUE"] != "NO APLICA")
    ].reset_index(drop = True)
    _datos = []
    _codes = _data["CODIGO OSINERGMIN"].unique()
    for _i in _codes:
        _data_code = _data[_data["CODIGO OSINERGMIN"] == _i]
        _tanks = _data_code["TANQUE"].unique()
        for _j in _tanks:
            _data_tank = _data_code[_data_code["TANQUE"] == _j]
            if (_data_tank['RESULTADO DE COMPARTIMIENTO'] == 'SIN FUGA').all():
                _result = "HERMETICO"
            else:
                _result = "NO HERMETICO"
            _datos.append(list(_data_tank.values[0]) + [_result, _data_tank["CAPACIDAD (GLN)"].sum()])
    _datos = pd.DataFrame(_datos, columns = [
        "CODIGO OSINERGMIN",  
        "TANQUE",
        "CAPACIDAD (GLN)",
        "RESULTADO DE COMPARTIMIENTO", 
        "EDAD DEL TANQUE",
        "RESULTADO",
        "CAPACIDAD DEL TANQUE (GLN)"
    ]).drop(columns = ["CAPACIDAD (GLN)", "RESULTADO DE COMPARTIMIENTO"])
    _datos = _datos[_datos["EDAD DEL TANQUE"] < 100]

    _hermeticos = alt.Chart(_datos[_datos["RESULTADO"] == "HERMETICO"]).mark_circle(
        size=50,
        stroke='black',
        strokeWidth=0.5
    ).encode(
        x=alt.X("EDAD DEL TANQUE:Q", title="AÑOS DE ANTIGUEDAD DEL TANQUE"),
        y=alt.Y("CAPACIDAD DEL TANQUE (GLN):Q", title="CAPACIDAD DEL TANQUE (GLN)"),
        color=alt.value("blue"),
        tooltip=[
            "CODIGO OSINERGMIN", "TANQUE", "CAPACIDAD DEL TANQUE (GLN)", "EDAD DEL TANQUE", "RESULTADO"
        ]
    )
    _no_hermeticos = alt.Chart(_datos[_datos["RESULTADO"] == "NO HERMETICO"]).mark_circle(
        size=50,
        stroke='black',
        strokeWidth=0.5
    ).encode(
        x="EDAD DEL TANQUE:Q",
        y="CAPACIDAD DEL TANQUE (GLN):Q",
        color=alt.value("red"),
        tooltip=[
            "CODIGO OSINERGMIN", "TANQUE", "CAPACIDAD DEL TANQUE (GLN)", "EDAD DEL TANQUE", "RESULTADO"
        ]
    )
    _chart = alt.layer(_hermeticos, _no_hermeticos).resolve_scale(
        color='independent'
    ).properties(
        title={
            "text":"RELACION ENTRE LOS AÑOS DE ANTIGUEDAD DEL TANQUE Y SU CAPACIDAD TOTAL, EN FUNCION DE SU HERMETICIDAD",
            "subtitle": _regiones,
            "subtitleFontSize": 5,
        },
    )
    mo.output.append(mo.ui.altair_chart(_chart))
    return


@app.cell
def _(Texto, matrix, mo):
    mo.output.append(mo.md("<br>"))
    download_3 = mo.ui.switch(label = "Exportar analisis")
    _texto = "ANALISIS CRÍTICO"
    mo.output.append(mo.hstack([Texto(_texto, 26, "left", "bold").create(), download_3], justify = "space-between", align = "stretch"))
    _texto = "En la gráfica, seleccione la barra correspondiente a la región competente que desea analizar, con el fin de verificar la cantidad de tanques no hermeticos con los que cuenta dicha región:"
    mo.output.append(Texto(_texto, 16, "justify", "normal").create())
    danger_data = matrix[
        ((matrix["RESULTADO DE COMPARTIMIENTO"] == "CON FUGA") | 
        (matrix["RESULTADO DE TUBERÍA"] == "CON FUGA")) & 
        (matrix["CONTENIDO"] != "GLP")
    ].reset_index(drop = True)
    return danger_data, download_3


@app.cell
def _(alt, danger_data, mo):
    _danger_counts = danger_data["OFICINA REGIONAL"].value_counts().reset_index()
    _danger_counts.columns = ["OFICINA REGIONAL", "ACREDITA FUGAS"]

    _danger_graph = alt.Chart(_danger_counts).mark_bar(color = "#ff1100").encode(
        x = alt.X("OFICINA REGIONAL", sort = '-y'),
        y = alt.Y("ACREDITA FUGAS")
    ).properties(
            title = "CANTIDAD DE COMPONENTES QUE ACREDITARON FUGAS EN LAS PRUEBAS EJECUTADAS"
        )
    danger_graph = mo.ui.altair_chart(_danger_graph)
    mo.output.append(danger_graph)
    return (danger_graph,)


@app.cell
def _(danger_data, danger_graph, download_3, io, mo):
    if danger_graph.value.empty:
        _data_filter = danger_data
    else:
        _data_filter = danger_data[danger_data["OFICINA REGIONAL"].isin(danger_graph.value["OFICINA REGIONAL"].tolist())].reset_index(drop = True)
    mo.output.append(_data_filter)

    if download_3.value:
        with mo.status.spinner(title = 'Generando Excel', subtitle = 'espere ...', remove_on_exit = True) as _spinner:
            _buffer = io.BytesIO()
            _data_filter.to_excel(_buffer, index = False, engine = "openpyxl")
            _buffer.seek(0)
            _excel_download = mo.download(
                data = _buffer.read(),
                filename = "AplicarMedidasdeSeguridad.xlsx",
                mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                label = "Descargar Excel"
            )
            mo.output.append(_excel_download)
    return


@app.cell
def _(Texto, mo):
    mo.output.append(mo.md("<br>"))
    mo.output.append(mo.md("---"))
    mo.output.append(mo.hstack([Texto("Documentacion del proyecto en: ", 12, "center", "normal").create(), mo.md("[![DeepWiki](https://img.shields.io/badge/DeepWiki-junior19a2000%2FSTE-blue.svg?logo=data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAACwAAAAyCAYAAAAnWDnqAAAAAXNSR0IArs4c6QAAA05JREFUaEPtmUtyEzEQhtWTQyQLHNak2AB7ZnyXZMEjXMGeK/AIi+QuHrMnbChYY7MIh8g01fJoopFb0uhhEqqcbWTp06/uv1saEDv4O3n3dV60RfP947Mm9/SQc0ICFQgzfc4CYZoTPAswgSJCCUJUnAAoRHOAUOcATwbmVLWdGoH//PB8mnKqScAhsD0kYP3j/Yt5LPQe2KvcXmGvRHcDnpxfL2zOYJ1mFwrryWTz0advv1Ut4CJgf5uhDuDj5eUcAUoahrdY/56ebRWeraTjMt/00Sh3UDtjgHtQNHwcRGOC98BJEAEymycmYcWwOprTgcB6VZ5JK5TAJ+fXGLBm3FDAmn6oPPjR4rKCAoJCal2eAiQp2x0vxTPB3ALO2CRkwmDy5WohzBDwSEFKRwPbknEggCPB/imwrycgxX2NzoMCHhPkDwqYMr9tRcP5qNrMZHkVnOjRMWwLCcr8ohBVb1OMjxLwGCvjTikrsBOiA6fNyCrm8V1rP93iVPpwaE+gO0SsWmPiXB+jikdf6SizrT5qKasx5j8ABbHpFTx+vFXp9EnYQmLx02h1QTTrl6eDqxLnGjporxl3NL3agEvXdT0WmEost648sQOYAeJS9Q7bfUVoMGnjo4AZdUMQku50McDcMWcBPvr0SzbTAFDfvJqwLzgxwATnCgnp4wDl6Aa+Ax283gghmj+vj7feE2KBBRMW3FzOpLOADl0Isb5587h/U4gGvkt5v60Z1VLG8BhYjbzRwyQZemwAd6cCR5/XFWLYZRIMpX39AR0tjaGGiGzLVyhse5C9RKC6ai42ppWPKiBagOvaYk8lO7DajerabOZP46Lby5wKjw1HCRx7p9sVMOWGzb/vA1hwiWc6jm3MvQDTogQkiqIhJV0nBQBTU+3okKCFDy9WwferkHjtxib7t3xIUQtHxnIwtx4mpg26/HfwVNVDb4oI9RHmx5WGelRVlrtiw43zboCLaxv46AZeB3IlTkwouebTr1y2NjSpHz68WNFjHvupy3q8TFn3Hos2IAk4Ju5dCo8B3wP7VPr/FGaKiG+T+v+TQqIrOqMTL1VdWV1DdmcbO8KXBz6esmYWYKPwDL5b5FA1a0hwapHiom0r/cKaoqr+27/XcrS5UwSMbQAAAABJRU5ErkJggg==)](https://deepwiki.com/junior19a2000/STE)")], justify = "center", align = "center"))
    return


if __name__ == "__main__":
    app.run()
